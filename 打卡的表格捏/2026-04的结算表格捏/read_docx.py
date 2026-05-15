"""
uv pip install python-docx
uv run read_docx.py
"""

from docx import Document

path = r".\软件2312-张喻-集美大学学生校内勤工助学工时记录表（记录表编号：JD02XS00807） (1).docx"

doc = Document(path)

print("=== 段落 ===")
for i, para in enumerate(doc.paragraphs):
    if para.text.strip():
        print(f"{i}: {para.text}")

print("\n=== 表格 ===")
for ti, table in enumerate(doc.tables):
    print(f"\n--- 表格 {ti} ---")
    for ri, row in enumerate(table.rows):
        cells = [cell.text.strip() for cell in row.cells]
        print(f"行{ri}: {cells}")
